# -*- coding: utf-8 -*-
"""Generator of bills in the format of Приложение №2 (amendments) and Приложение №1 (new law).

The title block is copied verbatim from the official templates (ACT. 000 (1).docx / (2).docx),
so fonts, horizontal rules, headers and footers stay identical to the template.
"""
import copy, difflib, re, os
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
TPL_AMEND = os.path.join(HERE, '..', 'ACT. 000 (1).docx')
TPL_NEW = os.path.join(HERE, '..', 'ACT. 000 (2).docx')

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
RED = 'F4CCCC'
GREEN = 'D9EAD3'


def _el(tag, **attrs):
    e = etree.SubElement(etree.Element('dummy'), qn(tag))
    e.getparent().remove(e)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def _rpr(bold=False, fill=None, size=24, color='333333', font='Times New Roman'):
    rpr = _el('w:rPr')
    f = _el('w:rFonts', **{'w:ascii': font, 'w:eastAsia': font, 'w:hAnsi': font, 'w:cs': font})
    rpr.append(f)
    if bold:
        rpr.append(_el('w:b'))
        rpr.append(_el('w:bCs'))
    if color:
        rpr.append(_el('w:color', **{'w:val': color}))
    rpr.append(_el('w:sz', **{'w:val': size}))
    rpr.append(_el('w:szCs', **{'w:val': size}))
    if fill:
        rpr.append(_el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': fill}))
    return rpr


def _run(text, bold=False, fill=None, size=24, color='333333'):
    r = _el('w:r')
    r.append(_rpr(bold=bold, fill=fill, size=size, color=color))
    t = _el('w:t')
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


def _para(runs, after=None, first_line=None, left=None, hanging=None, jc=None, cell=False):
    p = _el('w:p')
    ppr = _el('w:pPr')
    if cell:
        ppr.append(_el('w:widowControl', **{'w:val': '0'}))
        ppr.append(_el('w:spacing', **{'w:line': '240', 'w:lineRule': 'auto'}))
    elif after is not None:
        ppr.append(_el('w:spacing', **{'w:after': str(after)}))
    ind = {}
    if first_line is not None:
        ind['w:firstLine'] = str(first_line)
    if left is not None:
        ind['w:left'] = str(left)
    if hanging is not None:
        ind['w:hanging'] = str(hanging)
    if ind:
        ppr.append(_el('w:ind', **ind))
    if jc:
        ppr.append(_el('w:jc', **{'w:val': jc}))
    p.append(ppr)
    for r in runs:
        p.append(r)
    return p


# ----------------------------------------------------------------------------------
# word level diff
# ----------------------------------------------------------------------------------
_TOK = re.compile(r'\n|[ \t]+|[^\s]+')


def _tokens(text):
    return _TOK.findall(text)


def _smooth(marks, tokens, max_gap=3):
    """Absorb tiny unchanged islands (<= max_gap non-newline tokens, or pure whitespace)
    that are sandwiched between changed tokens, to keep highlights readable."""
    n = len(marks)
    i = 0
    while i < n:
        if not marks[i]:
            j = i
            while j < n and not marks[j]:
                j += 1
            island = tokens[i:j]
            if i > 0 and j < n and '\n' not in island:
                words = [t for t in island if t.strip()]
                if len(words) <= max_gap:
                    for k in range(i, j):
                        marks[k] = True
            i = j
        else:
            i += 1
    return marks


def diff_marks(old, new):
    a, b = _tokens(old), _tokens(new)
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    ma, mb = [False] * len(a), [False] * len(b)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == 'equal':
            continue
        for i in range(i1, i2):
            ma[i] = True
        for j in range(j1, j2):
            mb[j] = True
    ma = _smooth(ma, a)
    mb = _smooth(mb, b)
    return list(zip(a, ma)), list(zip(b, mb))


def _marked_paragraphs(marked, fill):
    """Turn a marked token stream into a list of paragraphs (each list of (text, changed))."""
    paras, cur = [], []
    for tok, ch in marked:
        if tok == '\n':
            paras.append(cur)
            cur = []
        else:
            cur.append((tok, ch))
    paras.append(cur)
    out = []
    for para in paras:
        # merge consecutive tokens with the same mark
        segs = []
        for tok, ch in para:
            if segs and segs[-1][1] == ch:
                segs[-1][0] += tok
            else:
                segs.append([tok, ch])
        # never shade leading/trailing pure whitespace
        runs = []
        for txt, ch in segs:
            if ch and not txt.strip():
                ch = False
            runs.append(_run(txt, fill=(fill if ch else None), color=None))
        out.append(_para(runs, cell=True))
    return out


# ----------------------------------------------------------------------------------
# table
# ----------------------------------------------------------------------------------

def _tcpr():
    tcpr = _el('w:tcPr')
    tcpr.append(_el('w:tcW', **{'w:w': '5940', 'w:type': 'dxa'}))
    tcpr.append(_el('w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': 'auto'}))
    mar = _el('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        mar.append(_el('w:' + side, **{'w:w': '100', 'w:type': 'dxa'}))
    tcpr.append(mar)
    return tcpr


def _cell(paragraphs):
    tc = _el('w:tc')
    tc.append(_tcpr())
    if not paragraphs:
        paragraphs = [_para([], cell=True)]
    for p in paragraphs:
        tc.append(p)
    return tc


def _table(rows):
    tbl = _el('w:tbl')
    tblpr = _el('w:tblPr')
    tblpr.append(_el('w:tblStyle', **{'w:val': 'a5'}))
    tblpr.append(_el('w:tblW', **{'w:w': '11880', 'w:type': 'dxa'}))
    tblpr.append(_el('w:tblInd', **{'w:w': '-1440', 'w:type': 'dxa'}))
    borders = _el('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(_el('w:' + side, **{'w:val': 'single', 'w:sz': '8', 'w:space': '0', 'w:color': '000000'}))
    tblpr.append(borders)
    tblpr.append(_el('w:tblLayout', **{'w:type': 'fixed'}))
    tblpr.append(_el('w:tblLook', **{'w:val': '0600', 'w:firstRow': '0', 'w:lastRow': '0', 'w:firstColumn': '0',
                                      'w:lastColumn': '0', 'w:noHBand': '1', 'w:noVBand': '1'}))
    tbl.append(tblpr)
    grid = _el('w:tblGrid')
    grid.append(_el('w:gridCol', **{'w:w': '5940'}))
    grid.append(_el('w:gridCol', **{'w:w': '5940'}))
    tbl.append(grid)
    # header row
    tr = _el('w:tr')
    tr.append(_cell([_para([_run('настоящая редакция', color=None)], cell=True, jc='center')]))
    tr.append(_cell([_para([_run('предложенный вариант', color=None)], cell=True, jc='center')]))
    tbl.append(tr)
    for old, new in rows:
        tr = _el('w:tr')
        if old is None:
            left = [_para([_run('Статья отсутствует.', color=None)], cell=True)]
            right = _marked_paragraphs([(t, True) for t in _tokens(new)], GREEN)
        elif new is None:
            left = _marked_paragraphs([(t, True) for t in _tokens(old)], RED)
            right = [_para([_run('Статья исключается.', color=None)], cell=True)]
        else:
            ma, mb = diff_marks(old, new)
            left = _marked_paragraphs(ma, RED)
            right = _marked_paragraphs(mb, GREEN)
        tr.append(_cell(left))
        tr.append(_cell(right))
        tbl.append(tr)
    return tbl


# ----------------------------------------------------------------------------------
# documents
# ----------------------------------------------------------------------------------

def _prepare(template, date_en, act_no=None):
    doc = Document(template)
    body = doc.element.body
    children = list(body)
    sectpr = children[-1]
    keep = children[:8]  # title block: ACT. 000, rule, senate line, date, rule, ЗАКОНОПРОЕКТ, subtitle, blank
    for ch in children[8:-1]:
        body.remove(ch)
    # replace the date text
    date_p = keep[3]
    ts = date_p.findall('.//' + qn('w:t'))
    ts[0].text = date_en
    for t in ts[1:]:
        t.text = ''
    if act_no:
        ts0 = keep[0].findall('.//' + qn('w:t'))
        ts0[0].text = act_no
        for t in ts0[1:]:
            t.text = ''
    return doc, body, sectpr


def _heading(text, after=300, left=851, hanging=425):
    return _para([_run(text, bold=True)], after=after, left=left, hanging=hanging)


def build_amendment(path, date_en, summary, sections, act_no=None):
    """sections: list of (law_name, rows) where rows = list of (old_text|None, new_text|None)."""
    doc, body, sectpr = _prepare(TPL_AMEND, date_en, act_no)
    sectpr.addprevious(_heading('SECTION 1. СУТЬ ПРАВКИ.'))
    if isinstance(summary, str):
        summary = [summary]
    for s in summary:
        sectpr.addprevious(_para([_run(s, bold=True)], after=300, first_line=480))
    n = 2
    for law_name, rows in sections:
        sectpr.addprevious(_para([_run('SECTION. %d. ВНЕСЕНИЕ ПРАВКИ В ' % n, bold=True), _run(law_name, bold=True)],
                                 after=300, first_line=426))
        sectpr.addprevious(_table(rows))
        sectpr.addprevious(_para([], after=300))
        n += 1
    doc.save(path)
    return path


def build_new_law(path, date_en, law_title, law_text, act_no=None):
    doc, body, sectpr = _prepare(TPL_NEW, date_en, act_no)
    sectpr.addprevious(_para([_run('SECTION. 1. ПРИНЯТИЕ НОВОГО ЗАКОНА', bold=True)], after=300, left=960, hanging=480))
    sectpr.addprevious(_para([_run('Принимается закон «%s»' % law_title)], after=300, first_line=425))
    for line in law_text.split('\n'):
        sectpr.addprevious(_para([_run(line)], after=0))
    doc.save(path)
    return path


# ----------------------------------------------------------------------------------
# explanation document (plain)
# ----------------------------------------------------------------------------------

def build_explanation(path, title, intro, items):
    """items: list of (heading, [paragraph, ...])"""
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = Document()
    st = d.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)
    for para in intro:
        d.add_paragraph(para)
    for heading, paras in items:
        p = d.add_paragraph()
        r = p.add_run(heading)
        r.bold = True
        for para in paras:
            if para.startswith('- '):
                d.add_paragraph(para[2:], style='List Bullet')
            else:
                d.add_paragraph(para)
    d.save(path)
    return path
