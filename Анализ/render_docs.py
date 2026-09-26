from pathlib import Path
from copy import deepcopy
import json,re,difflib
from collections import OrderedDict
from docx import Document
from docx.shared import Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.enum.text import WD_ALIGN_PARAGRAPH
R=Path(__file__).resolve().parents[1]
B=json.loads((R/'Анализ/проекты.json').read_text());D=json.loads((R/'Анализ/база.json').read_text())
T=Document(R/'Законы/ACT. 000 (1).docx');NT=Document(R/'Законы/ACT. 000 (2).docx')
TABLE=deepcopy(T.tables[1]._tbl)
TITLES={k:v['title'].strip().replace('“','«').replace('”','»') for k,v in D.items()}
TITLES['157']='Закон «О деятельности Секретной Службы Соединенных Штатов Америки в штате San-Andreas»'
TITLES['172']='Закон «О системе ордеров штата San-Andreas»'
TITLES['154']='Приложение № 1 к Уголовному кодексу штата San-Andreas «О закрытых, охраняемых и особо охраняемых территориях штата San-Andreas»'

def norm(s):
 s=s.replace('\u200b','').strip()
 s=re.sub(r'(?mi)^ч\.\s*(\d+(?:\.\d+)*)\.?\s*',lambda m:'Часть '+m[1]+'. ',s)
 s=re.sub(r'(?mi)^п\.\s*(\d+(?:\.\d+)*)\.?\s*',lambda m:'Пункт '+m[1]+'. ',s)
 s=re.sub(r'\bч\.\s*','части ',s,flags=re.I)
 s=re.sub(r'\bст\.\s*','статьи ',s,flags=re.I)
 s=re.sub(r'\bп\.\s*','пункта ',s)
 s=re.sub(r'\bГл\.\s*','главы ',s)
 s=re.sub(r'(?m)[ \t]+$','',s)
 s=re.sub(r'\n{3,}','\n\n',s)
 return s

def ptext(p,s):
 prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
 p.clear();r=p.add_run(s)
 if prop is not None:r._r.insert(0,prop)
 return p

def textp(doc,s,bold=False):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(6)
 p.paragraph_format.line_spacing=1
 r=p.add_run(s);r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor.from_string('333333');r.bold=bold
 return p

def tokens(s):return re.findall(r'\n|[^\S\n]+|[^\s]+',norm(s))
def highlighted(cell,ts,flags,color):
 cell.text='';p=cell.paragraphs[0]
 for t,f in zip(ts,flags):
  if t=='\n':p=cell.add_paragraph();continue
  p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
  r=p.add_run(t);r.font.name='Times New Roman';r.font.size=Pt(12);r.font.color.rgb=RGBColor.from_string(color if f else '333333')
 # Same single-spaced body typography as the supplied table.
 for p in cell.paragraphs:
  p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)

def compare(doc,c):
 tbl=deepcopy(TABLE);doc._body._body.insert(len(doc._body._body)-1,tbl);t=Table(tbl,doc._body)
 a=tokens(c['old']);z=tokens(c['new']);af=[False]*len(a);zf=[False]*len(z)
 for op,i,j,k,l in difflib.SequenceMatcher(None,a,z,autojunk=False).get_opcodes():
  if op!='equal':af[i:j]=[True]*(j-i);zf[k:l]=[True]*(l-k)
 highlighted(t.cell(1,0),a,af,'FF0000');highlighted(t.cell(1,1),z,zf,'008000')
 # Repeat header on page breaks; retain exactly two logical rows.
 pr=t.rows[0]._tr.get_or_add_trPr();e=OxmlElement('w:tblHeader');pr.append(e)
 for row in t.rows:
  for el in list(row._tr.get_or_add_trPr()):
   if el.tag in [qn('w:trHeight'),qn('w:cantSplit')]:el.getparent().remove(el)
 doc.add_paragraph()

def base(template,b):
 d=Document(template)
 ptext(d.paragraphs[0],f"ACT. {b['num']:03d}")
 ptext(d.paragraphs[3],'SEPTEMBER 27, 2026')
 d.core_properties.title=f"ACT. {b['num']:03d}. {b['title']}";d.core_properties.subject=b['purpose'];d.core_properties.author='';d.core_properties.last_modified_by=''
 return d

for b in B:
 if 'new_law' in b:
  d=base(R/'Законы/ACT. 000 (2).docx',b)
  ptext(d.paragraphs[9],'Принимается закон «О нормативных правовых актах и единых наименованиях государственных органов штата San-Andreas» в следующей редакции:')
  for line in b['new_law'].splitlines():textp(d,line,line.startswith(('Статья','Закон')))
 else:
  d=base(R/'Законы/ACT. 000 (1).docx',b)
  ptext(d.paragraphs[9],b['purpose'])
  body=d._body._body;start=d.paragraphs[10]._p
  remove=False
  for el in list(body):
   if el is start:remove=True
   if remove and el.tag!=qn('w:sectPr'):body.remove(el)
  groups=OrderedDict()
  for c in b['changes']:groups.setdefault(c['file'],[]).append(c)
  for sec,(k,cs) in enumerate(groups.items(),2):
   textp(d,f'SECTION. {sec}. ВНЕСЕНИЕ ПРАВКИ В {TITLES[k]}',True)
   for c in cs:compare(d,c)
  # Entry into force is normative content in the existing SECTION 1, not a new layout section.
  enact=' Изменения вступают в силу со дня официального опубликования настоящего Закона после его принятия и подписания в установленном порядке.'
  if b['num'] in [3,4,14]:enact=' Настоящие изменения вступают в силу одновременно с взаимосвязанными изменениями закона о государственной тайне, Процессуального кодекса, закона об адвокатуре и Уголовного кодекса, обеспечивающими специальный доступ защитника и единую ответственность за незаконное получение тайны, но не ранее официального опубликования всех указанных изменений после их принятия и подписания в установленном порядке.'
  ptext(d.paragraphs[9],b['purpose']+enact)
 path=R/'Законопроекты'/f"ACT. {b['num']:03d} — {b['title']}.docx";d.save(path);b['path']=str(path.relative_to(R))
 e=Document();e.styles['Normal'].font.name='Times New Roman';e.styles['Normal'].font.size=Pt(12)
 e.add_heading(f"Пояснения к ACT. {b['num']:03d}",0);e.add_paragraph(b['title']);e.add_paragraph(b['purpose'])
 e.add_paragraph('Источник: исключительно предоставленные файлы папки «Законы». Предлагаемые нормы не выдаются за действующую редакцию. Номер ACT является рабочим номером пакета, а не подтвержденным регистрационным номером Сената.')
 if b.get('dependencies'):e.add_heading('Связанные проекты',1);e.add_paragraph(b['dependencies'])
 if 'new_law' in b:e.add_paragraph(b['explanation'])
 for c in b['changes']:
  e.add_heading(f"{c['file']}.txt — {'статья ' if not c['article'].startswith('Глава') else ''}{c['article']}",1)
  e.add_paragraph('Что и почему: '+c['why']);e.add_paragraph('Пример: '+c['example'])
  # Precisely enumerate affected structural units, including punctuation-level or name harmonization.
  a=norm(c['old']).splitlines();z=norm(c['new']).splitlines();affected=[]
  for op,i,j,k,l in difflib.SequenceMatcher(None,a,z,autojunk=False).get_opcodes():
   if op=='equal':continue
   for line in z[k:l]:
    if line.strip():affected.append(line)
  if affected:
   e.add_paragraph('Измененные положения предложенной редакции:')
   for line in affected:e.add_paragraph(line)
 e.add_heading('Оформление и пределы проверки',1)
 e.add_paragraph('Для поправок использована форма приложения № 2: отдельная таблица из двух колонок и строго двух строк на каждую статью (заголовок и полный текст обеих редакций). Для нового закона использована форма приложения № 1. Сохраняются размеры страницы, поля, заголовки и оформление шаблонов. Красным показаны заменяемые фрагменты текущей редакции, зеленым — новые; неизмененный текст оставлен темным. Сокращенные обозначения частей, статей и пунктов развернуты в обеих колонках только редакционно, без изменения содержания. Главы закона 163.txt приведены целиком, поскольку исходник не разбит на статьи.')
 e.add_paragraph('Непредоставленные Конституция, закон о судебной системе, изображения границ и действующие ведомственные указы не реконструировались. Отсутствие таких файлов не доказывает отсутствие актов. Новые положения подлежат принятию уполномоченными субъектами; исходные файлы не изменены.')
 e.save(R/'Пояснения'/f"ACT. {b['num']:03d} — пояснения.docx")
(R/'Анализ/манифест.json').write_text(json.dumps(B,ensure_ascii=False,indent=2))
print('Rendered',len(B),'bills and explanations')
